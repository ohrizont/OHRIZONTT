import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook

file = "train.xlsx"

# Variables
indep_vars = [
    "Ram","screen_pixels","SSD_GB",
    "gpu_intel","gpu_nvidia",
    "gpu_tier_low","gpu_tier_mid",
    "gpu_pro","cpu_high","cpu_pro",
    "is_premium_brand","is_ultralight"
]

dep_var = "Price_euros"

# Leer datos
df = pd.read_excel(file, sheet_name="trans2_train")

X = df[indep_vars]
y = df[dep_var]

# Añadir constante
X_const = sm.add_constant(X)

# Ajustar modelo
model = sm.OLS(y, X_const).fit()

# Tabla de coeficientes y estadísticos
results = pd.DataFrame({
    "variable": model.params.index,
    "coef": model.params.values,
    "std_err": model.bse.values,
    "t": model.tvalues.values,
    "p_value": model.pvalues.values
}).sort_values("p_value")

# Estadísticos globales del modelo
stats = pd.DataFrame({
    "metric": [
        "R2", "R2_adj", "AIC", "BIC",
        "F_stat", "F_pvalue",
        "n_obs"
    ],
    "value": [
        model.rsquared,
        model.rsquared_adj,
        model.aic,
        model.bic,
        model.fvalue,
        model.f_pvalue,
        int(model.nobs)
    ]
})

# Determinar nombre de hoja correlativo
book = load_workbook(file)
base = "log"
sheet_name = base
i = 1
while sheet_name in book.sheetnames:
    sheet_name = f"{base}_{i}"
    i += 1

# Guardar en Excel
with pd.ExcelWriter(file, engine="openpyxl", mode="a", if_sheet_exists="new") as writer:
    results.to_excel(writer, sheet_name=sheet_name, index=False, startrow=0)
    stats.to_excel(writer, sheet_name=sheet_name, index=False, startrow=len(results)+3)
